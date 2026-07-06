"""Convert a CSV to a presentation-ready .xlsx (avoids Excel's pt-BR ';'
delimiter problem and comma-in-field issues).

    python src/csv_to_xlsx.py data/processed/normalization_compare.csv
    # -> data/processed/normalization_compare.xlsx
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WIDTHS = {"id": 8, "text": 55, "gloss_input": 45, "gloss_output_1": 45, "gloss_output_2": 45,
          "output": 42, "output_review": 42, "gloss_output_llm": 42, "gloss_output_hybrid": 42}


def convert(src: Path, dst: Path | None = None) -> Path:
    dst = dst or src.with_suffix(".xlsx")
    with src.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        raise SystemExit(f"empty: {src}")

    wb = Workbook()
    ws = wb.active
    ws.title = "compare"
    header = rows[0]
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=(r > 1 and header[c-1] != "id"))
            if r == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4472C4")
    for c, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS.get(name, 30)
    ws.freeze_panes = "A2"
    wb.save(dst)
    print(f"{src}  ->  {dst}  ({len(rows)-1} rows)")
    return dst


if __name__ == "__main__":
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "data/processed/normalization_compare.csv")
    convert(src, Path(sys.argv[2]) if len(sys.argv) > 2 else None)
